"""
Sistema de diseño compartido para reconstruir las plantillas de Tessera con
el rebrand 2026. Expone colores, estilos de texto y componentes (botón
píldora, tarjeta plana, línea "loop", cabecera/pie) como funciones sobre la
API de python-pptx, para que build_hc.py / build_hc_finance.py compongan
cada diapositiva declarativamente en vez de a base de XML a mano.

Reglas de marca aplicadas aquí (Manual de marca - Tessera.pdf y
Guia_uso_rebrand_Tessera_2026.pptx):
- Ailerons: solo en el logotipo (ver logo_picture()).
- Raleway: todo el texto (Light=título, SemiBold=subtítulo, Regular=cuerpo,
  Medium=etiqueta, en Teal y versalitas con tracking).
- Manrope: todos los números, sin excepción — ver add_number().
- Separador: guion corto "-", nunca punto medio ni guion largo.
- Fondos planos (blanco/Bruma/Navy) + degradado Teal->Navy SOLO en
  portadas/cierres/banners. Sin sombras, brillos ni más degradados que ese.
- Único recurso decorativo permitido: una línea fina con degradado de 2
  colores ("Loop") — ver loop_accent_line(). Nada de círculos difuminados.
"""
from __future__ import annotations

import io
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parent.parent.parent
ASSETS = Path(__file__).resolve().parent / "assets"
AILERONS_FONT = ROOT / "Ailerons-Typeface.otf"

# ---- Colores de marca --------------------------------------------------
TEAL = RGBColor(0x58, 0x75, 0x79)
NAVY = RGBColor(0x20, 0x20, 0x31)
NAVY_DEEP = RGBColor(0x16, 0x16, 0x1F)
TEAL_LIGHT = RGBColor(0x8A, 0xA3, 0xA5)
BRUMA = RGBColor(0xE6, 0xEC, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ---- Tipografía ---------------------------------------------------------
RALEWAY_LIGHT = "Raleway Light"
RALEWAY = "Raleway"
RALEWAY_MEDIUM = "Raleway Medium"
RALEWAY_SEMIBOLD = "Raleway SemiBold"
MANROPE = "Manrope"
MANROPE_EXTRALIGHT = "Manrope ExtraLight"
MANROPE_MEDIUM = "Manrope Medium"

# ---- Geometría de la diapositiva (16:9, en puntos) ----------------------
SLIDE_W = 960.0
SLIDE_H = 540.0
MARGIN = 60.0
HEADER_Y = 28.0
FOOTER_Y = 510.0


def _set_letter_spacing(run, spacing_pt: float) -> None:
    """Tracking entre letras — python-pptx no lo expone, se escribe el
    atributo `spc` (centésimas de punto) directamente en <a:rPr>."""
    rPr = run.font._rPr
    rPr.set("spc", str(round(spacing_pt * 100)))


def style_run(run, *, typeface=RALEWAY, size=14, color=NAVY, bold=False,
              spacing_pt=None, uppercase=False):
    if uppercase and run.text:
        run.text = run.text.upper()
    run.font.name = typeface
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if spacing_pt is not None:
        _set_letter_spacing(run, spacing_pt)
    return run


def add_textbox(slide, x, y, w, h, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
                 word_wrap=True, line_spacing=None):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = word_wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing is not None:
        p.line_spacing = line_spacing
    return box


def add_text(slide, x, y, w, h, text, *, typeface=RALEWAY, size=14, color=NAVY,
             bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=None, uppercase=False, spacing_pt=None):
    """Caja de texto de un solo párrafo/run — la forma más común: títulos,
    etiquetas, una frase de cuerpo. Devuelve la forma (shape) por si hace
    falta su shape_id para pptx_builder.py."""
    box = add_textbox(slide, x, y, w, h, align=align, anchor=anchor, line_spacing=line_spacing)
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    style_run(run, typeface=typeface, size=size, color=color, bold=bold,
              spacing_pt=spacing_pt, uppercase=uppercase)
    return box


def add_paragraphs(slide, x, y, w, h, paragraphs, *, typeface=RALEWAY, size=12,
                    color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
                    line_spacing=1.15, space_after=6, bold=False):
    """Varios párrafos independientes en una misma caja (p.ej. un cuerpo
    legal de términos y condiciones, o una bio con varias líneas)."""
    box = add_textbox(slide, x, y, w, h, align=align, anchor=anchor)
    tf = box.text_frame
    for i, text in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        style_run(run, typeface=typeface, size=size, color=color, bold=bold)
    return box


def add_number(slide, x, y, w, h, text, *, size=40, color=NAVY, weight="extralight",
                align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """Cualquier cifra (estadísticas, honorarios, meses de garantía, fechas,
    números de página...) va SIEMPRE en Manrope — nunca en Raleway."""
    typeface = {"extralight": MANROPE_EXTRALIGHT, "medium": MANROPE_MEDIUM}.get(weight, MANROPE)
    return add_text(slide, x, y, w, h, text, typeface=typeface, size=size, color=color,
                     align=align, anchor=anchor)


# ---- Fondos --------------------------------------------------------------
def set_solid_bg(slide, color: RGBColor):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def set_gradient_bg(slide, c1: RGBColor, c2: RGBColor, angle=45.0):
    """Único uso permitido de degradado: portadas, cierres y banners."""
    fill = slide.background.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    fill.gradient_angle = angle


# ---- Formas planas (tarjetas, líneas, botones) ---------------------------
def flat_rect(slide, x, y, w, h, fill: RGBColor | None, *, line_color=None, line_pt=0.75):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(line_pt)
    return shp


def thin_rule(slide, x, y, w, *, color=TEAL, height_pt=1.0):
    """Línea fina separadora — el único tipo de borde que usa la marca
    entre bloques/tarjetas."""
    return flat_rect(slide, x, y, w, height_pt, color)


def pill_button(slide, x, y, w, h, label, *, fill: RGBColor, text_color: RGBColor,
                 outline: RGBColor | None = None, size=13):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shp.shadow.inherit = False
    shp.adjustments[0] = 0.5  # radio = mitad del lado corto -> píldora
    if outline is not None:
        shp.fill.background()
        shp.line.color.rgb = outline
        shp.line.width = Pt(1.25)
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        shp.line.fill.background()
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Pt(18)
    tf.margin_top = tf.margin_bottom = Pt(4)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    style_run(run, typeface=RALEWAY_MEDIUM, size=size, color=text_color)
    return shp


def loop_accent_line(slide, x, y, w, *, c1=TEAL, c2=NAVY, height_pt=3.0, angle=0.0):
    """El único recurso decorativo aprobado por el manual de marca
    ("Loop"): una línea fina con degradado de 2 colores, para resaltar un
    título o cerrar visualmente un bloque. Nunca un fondo difuminado."""
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(height_pt))
    shp.shadow.inherit = False
    shp.line.fill.background()
    shp.fill.gradient()
    stops = shp.fill.gradient_stops
    stops[0].color.rgb = c1
    stops[0].position = 0.0
    stops[1].color.rgb = c2
    stops[1].position = 1.0
    shp.fill.gradient_angle = angle
    return shp


def photo_picture(slide, path, x, y, w, h):
    return slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w), Pt(h))


def cover_picture(slide, path, x, y, w, h):
    """Coloca una foto rellenando exactamente la caja w×h SIN deformarla —
    equivalente a `object-fit: cover`: recorta lo que sobre del lado más
    ancho en vez de estirar la imagen entera a la fuerza."""
    with Image.open(path) as im:
        iw, ih = im.size
    box_aspect = w / h
    img_aspect = iw / ih
    pic = slide.shapes.add_picture(str(path), Pt(x), Pt(y), Pt(w), Pt(h))
    if img_aspect > box_aspect:
        crop = 1 - (box_aspect / img_aspect)
        pic.crop_left = crop / 2
        pic.crop_right = crop / 2
    elif img_aspect < box_aspect:
        crop = 1 - (img_aspect / box_aspect)
        pic.crop_top = crop / 2
        pic.crop_bottom = crop / 2
    return pic


_logo_cache: dict[tuple, bytes] = {}


def logo_png_bytes(text="TESSERA", color=(32, 32, 49), size=48, width=420, height=70) -> bytes:
    """Genera el logotipo en Ailerons como PNG (igual que tessaix/assets.py
    en la app) — Ailerons nunca se usa como fuente nativa de PowerPoint."""
    key = (text, color, size, width, height)
    if key in _logo_cache:
        return _logo_cache[key]
    img = Image.new("RGBA", (width, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype(str(AILERONS_FONT), size)
    except Exception:
        font = ImageFont.load_default()
    bbox = draw.textbbox((0, 0), text, font=font)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.text(((width - tw) // 2, (height - th) // 2 - bbox[1]), text, font=font, fill=color)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    data = buf.getvalue()
    _logo_cache[key] = data
    return data


def logo_picture(slide, x, y, h_pt, *, dark_bg=False):
    """Logotipo TESSERA arriba a la izquierda — blanco sobre fondo oscuro,
    Navy sobre fondo claro."""
    color = (255, 255, 255) if dark_bg else (32, 32, 49)
    data = logo_png_bytes(color=color, width=560, height=96)
    aspect = 560 / 96
    w_pt = h_pt * aspect
    stream = io.BytesIO(data)
    return slide.shapes.add_picture(stream, Pt(x), Pt(y), Pt(w_pt), Pt(h_pt))


# ---- Cabecera / pie compartidos ------------------------------------------
def header_chrome(slide, eyebrow: str, *, dark_bg=False):
    """Logo arriba-izquierda + etiqueta de sección arriba-derecha — igual
    en todas las diapositivas de contenido (no en portada/cierre)."""
    logo_picture(slide, MARGIN, HEADER_Y, 14)
    color = TEAL_LIGHT if dark_bg else TEAL
    add_text(
        slide, SLIDE_W - MARGIN - 300, HEADER_Y, 300, 16, eyebrow,
        typeface=RALEWAY_MEDIUM, size=10, color=color, align=PP_ALIGN.RIGHT,
        uppercase=True, spacing_pt=1.2,
    )
    thin_rule(slide, MARGIN, HEADER_Y + 24, SLIDE_W - 2 * MARGIN,
              color=(NAVY_DEEP if dark_bg else BRUMA), height_pt=1.0)


def footer_chrome(slide, page_number: int, *, dark_bg=False):
    color = TEAL_LIGHT if dark_bg else TEAL
    num_color = WHITE if dark_bg else NAVY
    add_text(slide, MARGIN, FOOTER_Y, 300, 16, "tesseraservices.com",
             typeface=RALEWAY, size=9, color=color)
    add_number(slide, SLIDE_W - MARGIN - 60, FOOTER_Y - 6, 60, 26, f"{page_number:02d}",
               size=16, color=num_color, weight="medium", align=PP_ALIGN.RIGHT)


def new_slide(prs):
    """Diapositiva en blanco (layout 6 = Blank en la plantilla por defecto
    de python-pptx)."""
    return prs.slides.add_slide(prs.slide_layouts[6])
