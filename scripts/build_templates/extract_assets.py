"""
Extrae las imágenes de CONTENIDO real (fotos de equipo, logos de clientes)
de las plantillas actuales antes de reconstruirlas desde cero con el nuevo
branding. Todo lo demás visual (tarjetas, iconos decorativos, el gráfico de
onda) se descarta a propósito: se reconstruye en plano con brand_kit.py.

Uso: python scripts/build_templates/extract_assets.py
"""
import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent.parent
ASSETS = Path(__file__).resolve().parent / "assets"
ASSETS.mkdir(exist_ok=True)


def dump_pictures(pptx_path: Path, slide_index: int, prefix: str):
    """
    OJO: python-pptx's `shape.image` devuelve el blob ORIGINAL completo,
    sin aplicar el recorte (crop_left/top/right/bottom) que la propia
    diapositiva le aplicó al colocarlo — hay que recortarlo nosotros con
    PIL o se extrae la foto entera (p.ej. un cuerpo completo) en vez del
    encuadre de cabeza que realmente se veía en la plantilla.
    """
    prs = Presentation(str(pptx_path))
    slide = prs.slides[slide_index]
    print(f"\n=== {pptx_path.name} slide {slide_index + 1} ===")
    n = 0
    for shp in slide.shapes:
        if shp.shape_type != 13:  # MSO_SHAPE_TYPE.PICTURE
            continue
        n += 1
        w_pt = Emu(shp.width).pt if shp.width else 0
        h_pt = Emu(shp.height).pt if shp.height else 0
        x_pt = Emu(shp.left).pt if shp.left is not None else 0
        y_pt = Emu(shp.top).pt if shp.top is not None else 0
        try:
            image = shp.image
        except ValueError as e:
            print(f"  id={shp.shape_id:<4} name={shp.name!r:<30} "
                  f"x={x_pt:.0f} y={y_pt:.0f} w={w_pt:.0f} h={h_pt:.0f} -> SKIP ({e})")
            continue
        crop = (shp.crop_left, shp.crop_top, shp.crop_right, shp.crop_bottom)
        data = image.blob
        if any(crop):
            im = Image.open(io.BytesIO(data))
            iw, ih = im.size
            box = (
                round(crop[0] * iw), round(crop[1] * ih),
                round(iw - crop[2] * iw), round(ih - crop[3] * ih),
            )
            im = im.crop(box)
            buf = io.BytesIO()
            im.save(buf, "PNG")
            data = buf.getvalue()
            ext = "png"
        else:
            ext = image.ext
        out = ASSETS / f"{prefix}_{n:02d}_id{shp.shape_id}_{int(w_pt)}x{int(h_pt)}.{ext}"
        out.write_bytes(data)
        print(f"  id={shp.shape_id:<4} name={shp.name!r:<30} crop={crop} "
              f"x={x_pt:.0f} y={y_pt:.0f} w={w_pt:.0f} h={h_pt:.0f} -> {out.name}")


if __name__ == "__main__":
    ORIGINAL_HC = Path(__file__).resolve().parent / "assets" / "_original_plantilla_hc.pptx"
    hc_source = ORIGINAL_HC if ORIGINAL_HC.exists() else ROOT / "plantilla_hc.pptx"
    dump_pictures(hc_source, 1, "hc_equipo")
    dump_pictures(hc_source, 13, "hc_confian")
    dump_pictures(ROOT / "plantilla_hc_finance.pptx", 3, "hcf_equipo_sin_eduardo")
    dump_pictures(ROOT / "plantilla_hc_finance.pptx", 4, "hcf_equipo_con_eduardo")
    dump_pictures(ROOT / "plantilla_hc_finance.pptx", 12, "hcf_partners")
