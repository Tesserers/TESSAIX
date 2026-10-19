"""
Construye `plantilla_hc.pptx` (Human Capital solo) desde cero con el
sistema de diseño del rebrand 2026 (ver brand_kit.py). Mantiene exactamente
la misma información y estructura que la plantilla anterior (mismo orden
de diapositivas, mismos textos "old" que busca tessaix/pptx_builder.py,
mismas fotos de equipo y logos de clientes) — solo cambia la piel visual.

Uso: python scripts/build_templates/build_hc.py
Escribe sobre plantilla_hc.pptx en la raíz del repo (git-tracked, recuperable).
"""
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

import brand_kit as bk

ROOT = Path(__file__).resolve().parent.parent.parent
ASSETS = Path(__file__).resolve().parent / "assets"
OUT = ROOT / "plantilla_hc.pptx"

SHAPE_LOG: dict[str, dict] = {}


def log(slide_key, label, shape_or_value):
    val = shape_or_value.shape_id if hasattr(shape_or_value, "shape_id") else shape_or_value
    SHAPE_LOG.setdefault(slide_key, {})[label] = val


# ───────────────────────── 1. Portada ─────────────────────────────────
def slide_portada(prs):
    s = bk.new_slide(prs)
    bk.set_gradient_bg(s, bk.TEAL, bk.NAVY, angle=45)
    bk.logo_picture(s, bk.MARGIN, 60, 22, dark_bg=True)
    bk.add_text(s, bk.MARGIN, 150, 500, 18, "PROPUESTA COMERCIAL", typeface=bk.RALEWAY_MEDIUM,
                size=11, color=bk.TEAL_LIGHT, uppercase=True, spacing_pt=1.5)
    bk.add_text(s, bk.MARGIN, 180, 760, 150, "Queremos ser tu\npartner.",
                typeface=bk.RALEWAY_LIGHT, size=50, color=bk.WHITE)
    bk.add_text(s, bk.MARGIN, 330, 600, 40, "Better decisions, together.",
                typeface=bk.RALEWAY, size=15, color=bk.WHITE)
    bk.add_text(s, bk.MARGIN, 478, 400, 18, "tesseraservices.com",
                typeface=bk.RALEWAY, size=9, color=bk.TEAL_LIGHT)
    bk.add_number(s, bk.SLIDE_W - bk.MARGIN - 80, 478, 80, 18, "2026",
                  size=11, color=bk.TEAL_LIGHT, weight="medium", align=PP_ALIGN.RIGHT)
    return s


# ───────────────────────── 2. Equipo ──────────────────────────────────
def slide_equipo(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "NUESTRO EQUIPO")
    bk.add_text(s, bk.MARGIN, 70, 500, 60, "El Equipo", typeface=bk.RALEWAY_LIGHT, size=40, color=bk.NAVY)
    bk.add_number(s, bk.SLIDE_W - bk.MARGIN - 220, 86, 90, 36, "+21", size=26, color=bk.NAVY, weight="extralight", align=PP_ALIGN.RIGHT)
    bk.add_text(s, bk.SLIDE_W - bk.MARGIN - 120, 98, 120, 20, "consultores", typeface=bk.RALEWAY, size=11, color=bk.NAVY)

    team = [
        ("team_hc_eduardo_serrano.png", "EDUARDO SERRANO", "PARTNER",
         "17 años en el mundo del emprendimiento tanto como CFO de grandes startups "
         "internacionales como de CFO en fondos líderes de Venture Capital y Venture "
         "Builders, especializado en aportar valor al negocio y su expansión desde el "
         "área financiera, la gestión de tesorería y la consecución de financiación y "
         "capital para todas ellas, con experiencia en el sector publicitario como "
         "Business Global Director en Adglow."),
        ("team_hc_manuel_pina.png", "MANUEL PINA", "PARTNER",
         "Con más de 10 años de experiencia en Big Four en Auditoría, Due Diligence y "
         "M&A, posteriormente asumió posiciones como CFO en grandes pymes españolas. "
         "En la actualidad lidera Tessera Finance, coordinando un equipo especializado "
         "en servicios de CFO, transacciones, Due Diligence y asesoramiento financiero "
         "estratégico para compañías en distintas fases de crecimiento."),
        ("team_hc_edward_manrique.png", "EDWARD MANRIQUE", "HEAD OF TALENT",
         "Más de 8 años de experiencia en identificación y evaluación de talento. "
         "Experiencia en Seedtag, donde lideró la contratación de más de 200 "
         "profesionales en 16 países, así como en grandes entidades financieras "
         "internacionales como Santander, BNP Paribas y BBVA."),
    ]
    card_w, gap = 266.0, 21.0
    for i, (photo, name, role, bio) in enumerate(team):
        x = bk.MARGIN + i * (card_w + gap)
        photo_h = 140.0
        bk.cover_picture(s, ASSETS / photo, x, 140, card_w, photo_h)
        bk.loop_accent_line(s, x, 296, 60)
        bk.add_text(s, x, 308, card_w, 20, name, typeface=bk.RALEWAY_SEMIBOLD, size=13, color=bk.NAVY)
        bk.add_text(s, x, 328, card_w, 16, role, typeface=bk.RALEWAY_MEDIUM, size=9,
                    color=bk.TEAL, uppercase=True, spacing_pt=0.8)
        bk.add_paragraphs(s, x, 348, card_w, 155, [bio], size=9, color=bk.NAVY, line_spacing=1.18)

    bk.footer_chrome(s, 1)
    return s


# ───────────────── 3 / 4. Dónde estamos (todas / solo Madrid) ─────────
def _city_card(s, x, y, w, h, city, sub, *, big=False):
    bk.flat_rect(s, x, y, w, h, bk.BRUMA)
    bk.loop_accent_line(s, x + 20, y + 20, 50)
    size = 20 if big else 15
    bk.add_text(s, x + 20, y + 34, w - 40, 30, city, typeface=bk.RALEWAY_SEMIBOLD, size=size, color=bk.NAVY)
    bk.add_text(s, x + 20, y + 34 + size + 8, w - 40, 20, sub, typeface=bk.RALEWAY, size=10.5, color=bk.TEAL)


def _sedes_header(s):
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "DÓNDE ESTAMOS")
    bk.add_text(s, bk.MARGIN, 70, 700, 60, "Dónde estamos", typeface=bk.RALEWAY_LIGHT, size=40, color=bk.NAVY)
    bk.add_text(s, bk.MARGIN, 128, 760, 20,
                "El mismo estándar de servicio en cada oficina.", typeface=bk.RALEWAY, size=12, color=bk.NAVY)


def slide_sedes_todas(prs):
    s = bk.new_slide(prs)
    _sedes_header(s)
    card_w, gap = 266.0, 21.0
    cities = [("MADRID", "Sede central"), ("BILBAO", "País Vasco - España"), ("OVIEDO", "Asturias - España")]
    for i, (city, sub) in enumerate(cities):
        x = bk.MARGIN + i * (card_w + gap)
        _city_card(s, x, 170, card_w, 190, city, sub, big=True)
    bk.footer_chrome(s, 2)
    return s


def slide_sedes_madrid(prs):
    s = bk.new_slide(prs)
    _sedes_header(s)
    _city_card(s, bk.MARGIN, 170, 266, 190, "MADRID", "Sede central", big=True)
    bk.footer_chrome(s, 2)
    return s


# ───────────────── 5. Por qué Tessera (TAMBIÉN donante de Índice/Contexto) ─
PQT_EYEBROW = "POR QUÉ TESSERA"
PQT_TITLE = "Por qué Tessera"
PQT_SUBTITLE = "Procesos claros y compromisos que se cumplen: así movemos el talento."

WHY_CARDS = [
    ("Sin CVs al azar",
     "No enviamos el primer CV, buscamos a quien encaja de verdad. Candidatos filtrados sobre la mesa en 72 horas."),
    ("Sin pausas",
     "Tu servicio no parará."),
    ("Siempre contigo",
     "No desaparecemos tras la incorporación: cuidamos a la persona y al cliente."),
]
STRIP_TILES = [
    ("BOUTIQUE", "Equipo senior en cada búsqueda."),
    ("TRATO DIRECTO", "Sin intermediarios: trato de tú a tú."),
    ("VISIÓN DE NEGOCIO", "Estrategia diseñada para ti."),
]


def slide_por_que_tessera(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.BRUMA)
    bk.header_chrome(s, PQT_EYEBROW)
    title_shp = bk.add_text(s, bk.MARGIN, 70, 760, 60, PQT_TITLE, typeface=bk.RALEWAY_LIGHT, size=40, color=bk.NAVY)
    body_shp = bk.add_text(s, bk.MARGIN, 128, 760, 24, PQT_SUBTITLE, typeface=bk.RALEWAY, size=12, color=bk.NAVY)
    log("pqt", "title", title_shp)
    log("pqt", "body", body_shp)

    card_w, gap = 266.0, 21.0
    prune = []
    for i, (title, body) in enumerate(WHY_CARDS):
        x = bk.MARGIN + i * (card_w + gap)
        bg = bk.flat_rect(s, x, 172, card_w, 190, bk.WHITE)
        acc = bk.loop_accent_line(s, x + 20, 194, 50)
        box = bk.add_textbox(s, x + 20, 210, card_w - 40, 140)
        tf = box.text_frame
        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = title
        bk.style_run(r1, typeface=bk.RALEWAY_SEMIBOLD, size=14, color=bk.NAVY)
        p1.space_after = Pt(6)
        p2 = tf.add_paragraph()
        p2.line_spacing = 1.2
        r2 = p2.add_run()
        r2.text = body
        bk.style_run(r2, typeface=bk.RALEWAY, size=10.5, color=bk.NAVY)
        prune += [bg.shape_id, acc.shape_id, box.shape_id]

    strip_y = 392
    tile_w = (760 - 2 * 16) / 3
    for i, (label, text) in enumerate(STRIP_TILES):
        x = bk.MARGIN + i * (tile_w + 16)
        bg = bk.flat_rect(s, x, strip_y, tile_w, 70, bk.WHITE)
        lbl = bk.add_text(s, x + 16, strip_y + 14, tile_w - 32, 16, label, typeface=bk.RALEWAY_MEDIUM,
                          size=9, color=bk.TEAL, uppercase=True, spacing_pt=0.8)
        txt = bk.add_text(s, x + 16, strip_y + 34, tile_w - 32, 28, text, typeface=bk.RALEWAY, size=10, color=bk.NAVY)
        prune += [bg.shape_id, lbl.shape_id, txt.shape_id]

    log("pqt", "prune_ids", prune)
    bk.footer_chrome(s, 3)
    return s


# ───────────────── 6. Headhunting ──────────────────────────────────────
def _service_card(s, x, y, w, h, header, rows, *, header_color=bk.TEAL):
    bk.flat_rect(s, x, y, w, h, bk.BRUMA)
    bk.add_text(s, x + 18, y + 16, w - 36, 16, header, typeface=bk.RALEWAY_MEDIUM, size=10,
                color=header_color, uppercase=True, spacing_pt=0.8)
    bk.loop_accent_line(s, x + 18, y + 36, 44)
    box = bk.add_textbox(s, x + 18, y + 48, w - 36, h - 64)
    tf = box.text_frame
    first = True
    for title, body in rows:
        p1 = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p1.space_before = Pt(0 if p1 is tf.paragraphs[0] else 10)
        r1 = p1.add_run()
        r1.text = title
        bk.style_run(r1, typeface=bk.RALEWAY_SEMIBOLD, size=10.5, color=bk.NAVY)
        p2 = tf.add_paragraph()
        p2.line_spacing = 1.15
        r2 = p2.add_run()
        r2.text = body
        bk.style_run(r2, typeface=bk.RALEWAY, size=9.5, color=bk.NAVY)
    return box


def slide_headhunting(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "HEADHUNTING")
    bk.add_text(s, bk.MARGIN, 70, 700, 50, "Headhunting", typeface=bk.RALEWAY_LIGHT, size=36, color=bk.NAVY)

    card_w, gap = 266.0, 21.0
    box1 = _service_card(s, bk.MARGIN, 140, card_w, 270, "¿Por qué elegirnos?", [
        ("VELOCIDAD", "Disponemos de nuestra propia metodología, la cual nos permite presentar candidatos a tiempo."),
        ("ESPECIALIZACIÓN", "Conocemos tu sector y cómo se construyen sus equipos."),
        ("ÉXITO COMPARTIDO", "El equipo que lleva tu negocio al siguiente nivel."),
    ])
    box2 = _service_card(s, bk.MARGIN + card_w + gap, 140, card_w, 270, "Beneficios para ti", [
        ("CERO ESTRÉS", "Nos ocupamos de todo. Tú solo conoces a los mejores."),
        ("TALENTO A MEDIDA", "Candidatos que no solo encajan: impulsan tu crecimiento."),
        ("RELACIÓN A LARGO PLAZO", "Tu éxito es también el nuestro."),
    ])
    box3 = _service_card(s, bk.MARGIN + 2 * (card_w + gap), 140, card_w, 270, "¿Cómo lo hacemos?", [
        ("ELEGIMOS CON PRECISIÓN", "Habilidades, actitud y encaje cultural."),
        ("SOLO LO MEJOR", "Únicamente candidatos que realmente suman."),
        ("ACOMPAÑAMOS EL PROCESO", "Contigo también después de la incorporación."),
    ])
    log("headhunting", "card_why", box1)
    log("headhunting", "card_benefits", box2)
    log("headhunting", "card_how", box3)

    bk.flat_rect(s, bk.MARGIN, 428, 760, 76, bk.NAVY)
    bk.add_text(s, bk.MARGIN + 20, 442, 140, 16, "IA PROPIA DE SELECCIÓN", typeface=bk.RALEWAY_MEDIUM,
                size=9, color=bk.TEAL_LIGHT, uppercase=True, spacing_pt=0.8)
    bk.add_text(s, bk.MARGIN + 20, 460, 140, 30, "TESSA", typeface=bk.RALEWAY_SEMIBOLD, size=18, color=bk.WHITE)
    bk.add_text(s, bk.MARGIN + 170, 442, 420, 48,
                "Nuestra inteligencia artificial, junto a nuestra metodología propia, identifica "
                "y prioriza al mejor talento del mercado en cada búsqueda.",
                typeface=bk.RALEWAY, size=10, color=bk.WHITE, line_spacing=1.2)
    bk.add_number(s, bk.MARGIN + 610, 440, 90, 36, "72h", size=24, color=bk.WHITE, weight="extralight")
    bk.add_text(s, bk.MARGIN + 610, 476, 150, 24, "Primer shortlist\nsobre tu mesa",
                typeface=bk.RALEWAY, size=8.5, color=bk.TEAL_LIGHT, line_spacing=1.05)
    bk.footer_chrome(s, 4, dark_bg=False)
    return s


# ───────────────── 7. Outsourcing Time & Material ─────────────────────
def slide_outsourcing(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "OUTSOURCING TIME & MATERIAL")
    bk.add_text(s, bk.MARGIN, 70, 760, 50, "Outsourcing Time & Material",
                typeface=bk.RALEWAY_LIGHT, size=30, color=bk.NAVY)
    headline = bk.add_text(s, bk.MARGIN, 116, 760, 24,
                "Externalización que sí funciona: pagas por trabajo real, sin papeleos.",
                typeface=bk.RALEWAY_SEMIBOLD, size=14, color=bk.TEAL, line_spacing=1.15)
    body = bk.add_text(s, bk.MARGIN, 150, 760, 36,
                "Contratar cuesta más de lo que parece. Con el modelo Time & Material ajustas el "
                "equipo a tu actividad real y nosotros nos encargamos de toda la gestión.",
                typeface=bk.RALEWAY, size=11, color=bk.NAVY, line_spacing=1.2)
    log("outsourcing", "headline", headline)
    log("outsourcing", "body", body)

    cards = [
        ("Ahorro en costes de personal", "Pagas solo por horas reales de trabajo."),
        ("Cobertura total", "Cubrimos bajas y vacaciones."),
        ("Sin gestión", "Nóminas y trámites, a cargo nuestro."),
        ("Flexibilidad", "Escalas el equipo cuando quieras."),
    ]
    card_w, gap = 365.0, 20.0
    card_ids = []
    for i, (title, body_text) in enumerate(cards):
        col = i % 2
        row = i // 2
        x = bk.MARGIN + col * (card_w + gap)
        y = 230 + row * 80
        bg = bk.flat_rect(s, x, y, card_w, 68, bk.BRUMA)
        bk.add_text(s, x + 18, y + 12, card_w - 36, 18, title, typeface=bk.RALEWAY_SEMIBOLD, size=11.5, color=bk.NAVY)
        txt = bk.add_text(s, x + 18, y + 34, card_w - 36, 28, body_text, typeface=bk.RALEWAY, size=9.5, color=bk.NAVY)
        card_ids.append((bg.shape_id, txt.shape_id))
    log("outsourcing", "card1_body_shape", card_ids[0][1])
    log("outsourcing", "card2_body_shape", card_ids[1][1])
    log("outsourcing", "card3_body_shape", card_ids[2][1])
    log("outsourcing", "card4_body_shape", card_ids[3][1])

    badges = ["Recruiters especializados", "Cuidamos a las personas", "Partner de principio a fin"]
    badge_w = (760 - 2 * 16) / 3
    for i, label in enumerate(badges):
        x = bk.MARGIN + i * (badge_w + 16)
        bk.pill_button(s, x, 448, badge_w, 32, label, fill=bk.NAVY, text_color=bk.WHITE, size=10)
    bk.footer_chrome(s, 5)
    return s


# ───────────────── 8. Consultoría Salarial ─────────────────────────────
def slide_salarial(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "ESTRATEGIA SALARIAL")
    bk.add_text(s, bk.MARGIN, 70, 760, 50, "Consultoría Salarial", typeface=bk.RALEWAY_LIGHT, size=34, color=bk.NAVY)
    headline = bk.add_text(s, bk.MARGIN, 116, 760, 24, "Inteligencia retributiva para decisiones que importan.",
                typeface=bk.RALEWAY_SEMIBOLD, size=13, color=bk.TEAL)
    body = bk.add_text(s, bk.MARGIN, 144, 760, 36,
                "Acompañamos a organizaciones en el diseño de estructuras salariales coherentes, "
                "alineadas con el mercado y preparadas para el nuevo marco regulatorio europeo.",
                typeface=bk.RALEWAY, size=11, color=bk.NAVY, line_spacing=1.2)
    log("salarial", "headline", headline)
    log("salarial", "body", body)

    bk.add_text(s, bk.MARGIN, 196, 300, 16, "QUÉ OFRECEMOS", typeface=bk.RALEWAY_MEDIUM,
                size=9.5, color=bk.TEAL, uppercase=True, spacing_pt=0.8)

    items = [
        ("Benchmarking salarial", "Análisis de bandas por perfil, sector y geografía con datos contrastados de múltiples fuentes."),
        ("Transparencia retributiva", "Arquitectura salarial alineada con la Directiva EU 2023/970: equidad, coherencia y trazabilidad."),
        ("Estudio de mercado", "Informe de posicionamiento competitivo: percepción de candidatos, gaps y tendencias por rol."),
        ("Informes de proceso", "Expectativas salariales declaradas durante búsquedas activas, documentadas y estructuradas."),
        ("Diseño de bandas", "Estructuras P25/Mediana/P75 vinculadas a niveles de responsabilidad y ciclo de performance."),
        ("Asesoría People", "Apoyo en decisiones retributivas: ofertas, revisiones, retención y análisis de pay gap."),
    ]
    card_w, gap = 246.0, 16.0
    for i, (title, body_text) in enumerate(items):
        col = i % 3
        row = i // 3
        x = bk.MARGIN + col * (card_w + gap)
        y = 220 + row * 150
        bk.flat_rect(s, x, y, card_w, 136, bk.BRUMA)
        bk.loop_accent_line(s, x + 18, y + 16, 40)
        bk.add_text(s, x + 18, y + 30, card_w - 36, 18, title, typeface=bk.RALEWAY_SEMIBOLD, size=11.5, color=bk.NAVY)
        bk.add_text(s, x + 18, y + 52, card_w - 36, 76, body_text, typeface=bk.RALEWAY, size=9, color=bk.NAVY, line_spacing=1.2)
    bk.footer_chrome(s, 6)
    return s


# ───────────────── 9. Formación ─────────────────────────────────────────
def slide_formacion(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "FORMACIÓN EN RECURSOS HUMANOS")
    bk.add_text(s, bk.MARGIN, 70, 760, 50, "Formación", typeface=bk.RALEWAY_LIGHT, size=36, color=bk.NAVY)
    body = bk.add_text(s, bk.MARGIN, 118, 760, 24,
                "Potenciamos el talento interno: consultoría en formación orientada a procesos y resultados ágiles.",
                typeface=bk.RALEWAY, size=11, color=bk.NAVY)
    log("formacion", "body", body)

    cols = [
        ("ESPECIALIZACIÓN", ["Gestión del talento y selección avanzada", "Employee Experience y retención",
                              "Estrategia y procesos de RRHH", "Procesos automatizados", "IA aplicada a procesos de HR"]),
        ("METODOLOGÍA PRÁCTICA", ["Workshops dinámicos", "Casos reales aplicados a tu empresa",
                                   "Sprints formativos con entregables", "Evaluación y seguimiento continuo"]),
        ("RESULTADOS ÁGILES", ["Diagnóstico de procesos y skills", "Iteraciones cortas, feedback continuo",
                                "KPIs de aprendizaje medibles", "Mejora continua post-formación"]),
    ]
    card_w, gap = 266.0, 21.0
    for i, (header, bullets) in enumerate(cols):
        x = bk.MARGIN + i * (card_w + gap)
        bk.flat_rect(s, x, 158, card_w, 230, bk.BRUMA)
        bk.add_text(s, x + 18, 174, card_w - 36, 16, header, typeface=bk.RALEWAY_MEDIUM,
                    size=9.5, color=bk.TEAL, uppercase=True, spacing_pt=0.8)
        bk.loop_accent_line(s, x + 18, 194, 44)
        bk.add_paragraphs(s, x + 18, 208, card_w - 36, 170,
                          [f"-  {b}" for b in bullets], size=10, color=bk.NAVY, line_spacing=1.25, space_after=6)

    bk.flat_rect(s, bk.MARGIN, 410, 760, 90, bk.NAVY)
    bk.add_text(s, bk.MARGIN + 20, 432, 500, 50,
                "Diagnosticamos, diseñamos e implantamos programas formativos con metodología ágil. "
                "La formación no se queda en teoría: se mide en resultados.",
                typeface=bk.RALEWAY, size=10.5, color=bk.WHITE, line_spacing=1.25)
    bk.add_number(s, bk.MARGIN + 580, 424, 100, 40, "100%", size=26, color=bk.WHITE, weight="extralight")
    bk.add_text(s, bk.MARGIN + 580, 462, 150, 24, "Programas diseñados\na medida",
                typeface=bk.RALEWAY, size=8.5, color=bk.TEAL_LIGHT, line_spacing=1.05)
    bk.footer_chrome(s, 7)
    return s


# ───────────────── 10. Cross-sell Tessera Finance ──────────────────────
def slide_cross_sell(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.NAVY)
    bk.header_chrome(s, "TESSERA FINANCE", dark_bg=True)
    bk.add_text(s, bk.MARGIN, 70, 300, 16, "EXPERTOS EN LAS EMPRESAS", typeface=bk.RALEWAY_MEDIUM,
                size=9.5, color=bk.TEAL_LIGHT, uppercase=True, spacing_pt=0.8)
    bk.add_text(s, bk.MARGIN, 92, 420, 90,
                "En Tessera, potenciamos el talento humano con la tecnología para transformar "
                "empresas. Somos una firma de consultoría estratégica especializada en Finanzas, "
                "donde ofrecemos entre otros los siguientes servicios:",
                typeface=bk.RALEWAY, size=11, color=bk.WHITE, line_spacing=1.25)

    services = ["Búsqueda de Financiación", "M&A", "Business Plan", "CFO Part-Time", "Due Diligence", "Restructuring"]
    bk.add_paragraphs(s, bk.MARGIN, 190, 420, 200, [f"-  {sv}" for sv in services],
                      typeface=bk.RALEWAY, size=12, color=bk.WHITE, line_spacing=1.6, space_after=4)

    panel_x = bk.MARGIN + 460
    bk.flat_rect(s, panel_x, 70, 300, 340, bk.NAVY_DEEP)
    bk.add_text(s, panel_x + 24, 96, 252, 40, "ASESORÍA LABORAL,\nFISCAL & CONTABLE",
                typeface=bk.RALEWAY_SEMIBOLD, size=14, color=bk.WHITE, line_spacing=1.15)
    bk.add_text(s, panel_x + 24, 150, 252, 160,
                "Impulsamos a tu empresa para que te centres en lo importante: tu negocio. Nos "
                "encargamos de que, a nivel fiscal, laboral y contable, no haya ningún problema, "
                "se optimice la gestión fiscal, se aprovechen las subvenciones laborales y se "
                "disponga de información contable veraz, al día y precisa para la toma de decisiones.",
                typeface=bk.RALEWAY, size=9.5, color=bk.TEAL_LIGHT, line_spacing=1.3)
    bk.loop_accent_line(s, panel_x + 24, 326, 60)
    bk.add_number(s, panel_x + 24, 338, 90, 34, "37%", size=24, color=bk.WHITE, weight="extralight")
    bk.add_text(s, panel_x + 24, 372, 252, 28, "Ahorro medio de costes en el análisis gratuito",
                typeface=bk.RALEWAY, size=8.5, color=bk.TEAL_LIGHT, line_spacing=1.1)
    bk.footer_chrome(s, 8, dark_bg=True)
    return s


# ───────────────── 11 / 12. Términos y condiciones ─────────────────────
WARRANTY_OLD_SENTENCE = (
    "Si el candidato o el cliente rescinde el contrato de trabajo dentro de los 3 meses "
    "siguientes a la fecha de contratación, Tessera Human Capital se compromete a "
    "proporcionar otro candidato y/o a activar otro proceso de selección, a elección del "
    "cliente, con el fin de sustituir al candidato sin coste adicional."
)
DATE_PLACEHOLDER = "{dd/mm/aaaa}"


def _terms_header(s, page_no):
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "TÉRMINOS Y CONDICIONES")
    bk.add_text(s, bk.MARGIN, 70, 700, 40, "Términos y condiciones", typeface=bk.RALEWAY_LIGHT, size=28, color=bk.NAVY)
    bk.footer_chrome(s, page_no)


def slide_terminos_1(prs):
    s = bk.new_slide(prs)
    _terms_header(s, 9)
    bk.add_text(s, bk.MARGIN, 112, 700, 16, "CONDICIONES SUBYACENTES DE LA PROPUESTA",
                typeface=bk.RALEWAY_MEDIUM, size=9.5, color=bk.TEAL, uppercase=True, spacing_pt=0.8)

    col_w = 360.0
    left = [
        "Las presentes condiciones generales se aplican a cualquier servicio profesional "
        "prestado entre las partes y son adicionales a cualquier acuerdo específico "
        "suscrito, el cual prevalecerá sobre estas en caso de discrepancia. En este "
        "documento: (i) «Tessera» significa Tessera Human Capital, S.L.; (ii) «Cliente» es "
        "el destinatario de la Propuesta y (iii) «Servicios» son los servicios "
        "profesionales cubiertos por la Propuesta.",
        "1. CONFIDENCIALIDAD",
        "Toda la información relativa a la empresa cliente, sus miembros y los candidatos "
        "presentados o contactados será tratada por ambas partes con la máxima "
        "confidencialidad, manteniendo la obligación de secreto profesional sobre toda "
        "dicha información y, en particular, sobre los datos personales de los candidatos, "
        "incluso una vez finalizado este acuerdo, todo ello de conformidad con la "
        "legislación española.",
        "Ambas partes se comprometen mutuamente a mantener la confidencialidad y a no "
        "revelar a terceros ninguna información confidencial facilitada por la otra parte, "
        "durante la vigencia de este acuerdo y con posterioridad al mismo. Este compromiso "
        "no impedirá a Tessera Human Capital hacer referencia a los servicios prestados al "
        "Cliente y al propio Cliente como parte de su cartera de clientes, con fines "
        "publicitarios y de marketing.",
        "2. CANDIDATURAS",
        "2.1 Todos los candidatos identificados y presentados por Tessera Human Capital al "
        "cliente quedarán reservados exclusivamente para la presente selección. Si alguno "
        "de dichos candidatos fuera contratado por el cliente en un puesto correspondiente "
        "al encargo, o en otro puesto, o en otra empresa del grupo, dentro de un plazo de "
        "12 meses tras la presentación del candidato, el cliente deberá abonar a Tessera "
        "Human Capital el importe previsto para la contratación inicial, o la totalidad de "
        "los honorarios calculados sobre la retribución bruta anual, incluidas las "
        "retribuciones en especie.",
        "2.2 Esta cláusula se aplica en los mismos términos si el cliente presenta a alguno "
        "de estos candidatos a otra persona física o jurídica y si el candidato se "
        "compromete con dicha persona.",
    ]
    right = [
        "3. GARANTÍA",
        WARRANTY_OLD_SENTENCE,
        "Esta garantía podrá aplicarse salvo en los siguientes casos:",
        "3.1 Si la rescisión del contrato de trabajo se debe a un exceso de plantilla o a "
        "un cambio geográfico inicialmente no previsto.",
        "3.2 Si las condiciones laborales del candidato contratado difieren "
        "significativamente de las inicialmente definidas en el contrato de trabajo.",
        "3.3 Si la rescisión del contrato de trabajo se debe a una falta grave por parte de "
        "la empresa contratante.",
        "3.4 Si no se han abonado todas las facturas adeudadas por el cliente dentro de los "
        "plazos establecidos y conforme a las condiciones generales.",
        "Si el candidato abandona la empresa transcurridos 6 meses desde la fecha de "
        "contratación (periodo de prueba), Tessera Human Capital no se hace responsable y, "
        "por tanto, no tiene obligación de cubrir el puesto.",
    ]

    def render_column(x, items):
        box = bk.add_textbox(s, x, 140, col_w, 350)
        tf = box.text_frame
        for i, text in enumerate(items):
            is_heading = text[0].isdigit() and ". " in text[:4] or text.startswith(("1.", "2.", "3."))
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = 1.08
            p.space_after = Pt(4)
            run = p.add_run()
            run.text = text
            if is_heading:
                bk.style_run(run, typeface=bk.RALEWAY_SEMIBOLD, size=9, color=bk.TEAL)
            else:
                bk.style_run(run, typeface=bk.RALEWAY, size=7.3, color=bk.NAVY)
        return box

    shp_l = render_column(bk.MARGIN, left)
    shp_r = render_column(bk.MARGIN + col_w + 40, right)
    log("terminos1", "left", shp_l)
    log("terminos1", "right", shp_r)
    return s


def slide_terminos_2(prs):
    s = bk.new_slide(prs)
    _terms_header(s, 10)

    col_w = 360.0
    left = [
        ".4. RESPONSABILIDAD",
        "Las siguientes aclaraciones legales no excluyen nuestras condiciones de garantía, "
        "que permanecen válidas conforme a los términos definidos en el apartado 3.",
        "4.1 Tessera Human Capital se compromete a prestar un servicio de calidad y "
        "garantiza, como mínimo, la devolución de las cantidades abonadas en aquellos casos "
        "en los que hubiera advertido al cliente de un incumplimiento grave y directo de "
        "sus obligaciones.",
        "4.2 Sin perjuicio de lo anterior, salvo en caso de dolo, Tessera Human Capital no "
        "asume responsabilidad alguna, ni tampoco su personal, por cualquier negligencia o "
        "incapacidad, mala conducta o práctica inadecuada o cualquier otra reclamación "
        "atribuida a los candidatos presentados.",
        ".5. PROTECCIÓN DE DATOS",
        "Ambas partes reconocen que cumplen con la normativa de protección de datos "
        "personales y han sido informadas de las finalidades del tratamiento de los datos "
        "facilitados en esta oferta comercial, basadas en el interés legítimo de ambas "
        "partes. Cualquier cuestión relativa al tratamiento de datos podrá dirigirse a "
        "Tessera Human Capital.",
        "La cesión de los CV se realiza exclusivamente en el marco de la selección "
        "contratada; una vez finalizada, el cliente se compromete a devolverlos o "
        "destruirlos, a la confidencialidad de los datos personales facilitados, a su no "
        "utilización fuera de la relación contractual y al cumplimiento del RGPD, siendo el "
        "único responsable de cualquier incumplimiento de dicha normativa.",
        "Tessera Human Capital, por su parte, garantiza que cuenta con el consentimiento de "
        "los candidatos para la cesión de sus datos a sus clientes durante los procesos de "
        "selección de personal.",
    ]
    right_top = [
        "6. LEY APLICABLE Y JURISDICCIÓN",
        "Este acuerdo se regirá e interpretará de conformidad con la legislación española. "
        "Para la resolución de cualquier controversia derivada de la interpretación o "
        "ejecución de este acuerdo, las partes acuerdan someterse a la jurisdicción de los "
        "Juzgados y Tribunales de Madrid, renunciando expresamente a cualquier otro fuero "
        "que pudiera corresponderles.",
        "7. VIGENCIA DE ESTE ACUERDO",
        "Este acuerdo entrará en vigor desde la fecha de su firma, con fecha de "
        "finalización en mayo de 2026.",
        "8. FIRMA DEL ACUERDO DE COLABORACIÓN",
        f"En Madrid, a {DATE_PLACEHOLDER}",
    ]

    box_l = bk.add_textbox(s, bk.MARGIN, 140, col_w, 330)
    tf = box_l.text_frame
    for i, text in enumerate(left):
        is_heading = text.startswith((".4", ".5"))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.08
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = text
        bk.style_run(run, typeface=bk.RALEWAY_SEMIBOLD if is_heading else bk.RALEWAY,
                    size=9 if is_heading else 7.3, color=bk.TEAL if is_heading else bk.NAVY)
    log("terminos2", "left", box_l)

    rx = bk.MARGIN + col_w + 40
    box_r = bk.add_textbox(s, rx, 140, col_w, 170)
    tf = box_r.text_frame
    for i, text in enumerate(right_top):
        is_heading = text.startswith(("6.", "7.", "8."))
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.08
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = text
        bk.style_run(run, typeface=bk.RALEWAY_SEMIBOLD if is_heading else bk.RALEWAY,
                    size=9 if is_heading else 7.3, color=bk.TEAL if is_heading else bk.NAVY)
    log("terminos2", "right", box_r)

    sig_y = 330
    bk.thin_rule(s, rx, sig_y, col_w, color=bk.BRUMA, height_pt=1.0)
    bk.add_text(s, rx, sig_y + 12, 200, 16, "FIRMA DE TESSERA", typeface=bk.RALEWAY_MEDIUM,
                size=9.5, color=bk.TEAL, uppercase=True, spacing_pt=0.8)
    firma_shp = bk.add_paragraphs(s, rx, sig_y + 34, col_w, 70,
                      ["Nombre: Manuel García Pina", f"Fecha: {DATE_PLACEHOLDER}", "Firma:"],
                      size=9.5, color=bk.NAVY, line_spacing=1.3, space_after=4)
    log("terminos2", "firma_tessera", firma_shp)
    bk.add_text(s, rx, sig_y + 120, 200, 16, "FIRMA DEL CLIENTE", typeface=bk.RALEWAY_MEDIUM,
                size=9.5, color=bk.TEAL, uppercase=True, spacing_pt=0.8)
    bk.thin_rule(s, rx, sig_y + 170, col_w, color=bk.BRUMA, height_pt=1.0)
    return s


# ───────────────── 13. Honorarios ───────────────────────────────────────
FEE_OLD = "13%"


def slide_honorarios(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "HONORARIOS")
    bk.add_text(s, bk.MARGIN, 70, 760, 40, "El valor que aportamos a tu proyecto",
                typeface=bk.RALEWAY_LIGHT, size=24, color=bk.NAVY)

    cards = [
        ("01", "Experiencia Directa",
         "Entendemos tu entorno, lo que nos permite analizar con precisión las dinámicas "
         "de negocio y los aspectos críticos que impactan en tu crecimiento."),
        ("02", "Enfoque Práctico",
         "Orientados a resultados, facilitando decisiones claras y ejecutables que "
         "optimizan procesos."),
        ("03", "Impacto Real",
         "Soluciones que impulsan eficiencias y mejoras sostenibles en el negocio de "
         "nuestros clientes."),
    ]
    card_w, gap = 246.0, 16.0
    for i, (num, title, body) in enumerate(cards):
        x = bk.MARGIN + i * (card_w + gap)
        bk.flat_rect(s, x, 130, card_w, 150, bk.BRUMA)
        bk.add_number(s, x + 18, 142, 60, 30, num, size=20, color=bk.TEAL, weight="medium")
        bk.add_text(s, x + 18, 176, card_w - 36, 20, title, typeface=bk.RALEWAY_SEMIBOLD, size=12, color=bk.NAVY)
        bk.add_text(s, x + 18, 198, card_w - 36, 72, body, typeface=bk.RALEWAY, size=9, color=bk.NAVY, line_spacing=1.2)

    fee_y = 310
    bk.flat_rect(s, bk.MARGIN, fee_y, 760, 120, bk.NAVY)
    bk.add_text(s, bk.MARGIN + 24, fee_y + 20, 300, 16, "HONORARIOS", typeface=bk.RALEWAY_MEDIUM,
                size=9.5, color=bk.TEAL_LIGHT, uppercase=True, spacing_pt=0.8)
    bk.add_text(s, bk.MARGIN + 24, fee_y + 42, 400, 24, "Modalidad a éxito en Headhunting",
                typeface=bk.RALEWAY_SEMIBOLD, size=13, color=bk.WHITE)
    bk.add_text(s, bk.MARGIN + 24, fee_y + 72, 300, 18, "Salario Bruto Anual",
                typeface=bk.RALEWAY, size=10, color=bk.TEAL_LIGHT)
    fee_shape = bk.add_number(s, bk.MARGIN + 640, fee_y + 30, 90, 50, FEE_OLD, size=34,
                              color=bk.WHITE, weight="extralight", align=PP_ALIGN.RIGHT)
    log("honorarios", "fee", fee_shape)
    bk.footer_chrome(s, 11, dark_bg=False)
    return s


# ───────────────── 14. Confían en nosotros ──────────────────────────────
def slide_confian(prs):
    s = bk.new_slide(prs)
    bk.set_solid_bg(s, bk.WHITE)
    bk.header_chrome(s, "CONFÍAN EN NOSOTROS")
    bk.add_text(s, bk.MARGIN, 70, 700, 50, "Confían en nosotros", typeface=bk.RALEWAY_LIGHT, size=34, color=bk.NAVY)

    logos = sorted((ASSETS / "logos").glob("logo_*.png"))
    cols, rows = 6, 3
    cell_w = (bk.SLIDE_W - 2 * bk.MARGIN) / cols
    cell_h = 90.0
    start_y = 150.0
    for i, logo_path in enumerate(logos[: cols * rows]):
        col, row = i % cols, i // cols
        x = bk.MARGIN + col * cell_w
        y = start_y + row * cell_h
        from PIL import Image
        with Image.open(logo_path) as im:
            iw, ih = im.size
        max_w, max_h = cell_w - 24, 44.0
        scale = min(max_w / iw, max_h / ih)
        w, h = iw * scale, ih * scale
        bk.photo_picture(s, logo_path, x + (cell_w - w) / 2, y + (cell_h - h) / 2, w, h)
    bk.footer_chrome(s, 12)
    return s


# ───────────────── 15. Cierre ───────────────────────────────────────────
def slide_cierre(prs):
    s = bk.new_slide(prs)
    bk.set_gradient_bg(s, bk.TEAL, bk.NAVY, angle=45)
    bk.logo_picture(s, bk.MARGIN, 60, 22, dark_bg=True)
    bk.add_text(s, bk.MARGIN, 200, 700, 60, "BUILDING BETTER TEAMS, TOGETHER",
                typeface=bk.RALEWAY_LIGHT, size=30, color=bk.WHITE)
    bk.add_text(s, bk.MARGIN, 300, 300, 18, "Email", typeface=bk.RALEWAY_MEDIUM, size=9.5,
                color=bk.TEAL_LIGHT, uppercase=True, spacing_pt=0.8)
    bk.add_text(s, bk.MARGIN, 318, 400, 20, "edward@tesseraservices.com", typeface=bk.RALEWAY, size=12, color=bk.WHITE)
    bk.add_text(s, bk.MARGIN, 350, 300, 18, "LinkedIn", typeface=bk.RALEWAY_MEDIUM, size=9.5,
                color=bk.TEAL_LIGHT, uppercase=True, spacing_pt=0.8)
    bk.add_text(s, bk.MARGIN, 368, 400, 20, "@tesseraservices", typeface=bk.RALEWAY, size=12, color=bk.WHITE)
    bk.add_text(s, bk.MARGIN, 478, 400, 18, "tesseraservices.com", typeface=bk.RALEWAY, size=9, color=bk.TEAL_LIGHT)
    return s


def build():
    prs = Presentation()
    prs.slide_width = Pt(bk.SLIDE_W)
    prs.slide_height = Pt(bk.SLIDE_H)

    slide_portada(prs)
    slide_equipo(prs)
    slide_sedes_todas(prs)
    slide_sedes_madrid(prs)
    slide_por_que_tessera(prs)
    slide_headhunting(prs)
    slide_outsourcing(prs)
    slide_salarial(prs)
    slide_formacion(prs)
    slide_cross_sell(prs)
    slide_terminos_1(prs)
    slide_terminos_2(prs)
    slide_honorarios(prs)
    slide_confian(prs)
    slide_cierre(prs)

    prs.save(str(OUT))
    print(f"Guardado {OUT} con {len(prs.slides)} diapositivas")
    for slide_key, shapes in SHAPE_LOG.items():
        print(f"  [{slide_key}]")
        for label, val in shapes.items():
            print(f"    {label}: {val}")


if __name__ == "__main__":
    build()
