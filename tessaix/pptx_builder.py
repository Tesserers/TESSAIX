"""Ensamblado del .pptx a partir de las plantillas de Tessera.

Hay dos plantillas maestras, una por línea de negocio (ver config.py):
`PLANTILLA_HC` (Human Capital solo) y `PLANTILLA_HC_FINANCE` (Human Capital
+ Finance). Cada una es un SUPERCONJUNTO: trae ya dentro los bloques
opcionales (equipo con/sin Eduardo Serrano, sedes Madrid-sola o las 3,
términos y condiciones + honorarios) como diapositivas o formas completas,
y este módulo decide en tiempo de generación cuáles conservar y cuáles
quitar, en vez de tener que maquetar nada desde cero.

Las diapositivas de "Índice" y "Contexto" son la excepción: no vienen en
la plantilla, así que se CLONAN a partir de una diapositiva ya existente
que tiene el "cascarón" correcto (logo, línea, texto de sección, título,
subtítulo, número de página) — se le quitan las tarjetas propias de esa
diapositiva y se reutiliza el subtítulo como cuerpo de texto, dejando que
el motor de ajuste (text_fit) lo redimensione. Así no hay que maquetar
nada nuevo a mano ni arriesgarse a que no case con el estilo del resto.
"""
from __future__ import annotations

import datetime
import io
import re
import shutil
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation

from .ai_content import generate_context_paragraph
from .config import BUSINESS_LINE_HC_FINANCE, PLANTILLA_HC, PLANTILLA_HC_FINANCE, SERVICES_HC
from .team import LOCATIONS_MADRID_ONLY
from .xml_utils import (
    enable_shrink_autofit,
    remove_content_type_override,
    remove_relationship,
    remove_shape_by_id,
    remove_sld_id,
    replace_text_and_fit,
    replace_text_in_xml,
    resolve_rid_for_target,
)

# ─── SLIDES DE CADA PLANTILLA (numeración fija, 1-indexada) ───────────────
# plantilla_hc.pptx (15 diapositivas antes de añadir Índice/Contexto):
#  1 portada · 2 equipo (fijo) · 3 sedes (3 ciudades) · 4 sedes (solo Madrid)
#  5 por qué Tessera · 6 headhunting · 7 outsourcing · 8 salarial · 9 formación
#  10 Tessera Services (estático) · 11-12 términos y condiciones · 13 honorarios
#  14 confían en nosotros · 15 cierre
HC_SLIDES_DYNAMIC_TEXT = ["slide5.xml", "slide6.xml", "slide7.xml", "slide8.xml", "slide9.xml"]
HC_LOCATION_SLIDES = {"all": "slide3.xml", "madrid": "slide4.xml"}
HC_TERMS_SLIDES = ["slide11.xml", "slide12.xml"]
HC_FEE_SLIDE = "slide13.xml"
HC_SERVICE_SLIDES = {"headhunting": "slide6.xml", "outsourcing": "slide7.xml",
                      "salary": "slide8.xml", "formacion": "slide9.xml"}

# plantilla_hc_finance.pptx (15 diapositivas antes de añadir Índice/Contexto):
#  1 portada · 2/3 "Tessera en una página" (3 sedes / solo Madrid, se
#  conserva una y se borra la otra) · 4/5 equipo (sin/con Eduardo Serrano,
#  igual: se conserva una) · 6 para el fondo · 7 servicios Finance
#  8 servicios Human Capital (resumen) · 9 por qué Tessera
#  10-11 términos y condiciones · 12 honorarios · 13 confían en nosotros
#  14 contacto · 15 cierre
HCF_LOCATION_SLIDES = {"all": "slide2.xml", "madrid": "slide3.xml"}
HCF_TEAM_SLIDES = {"no_eduardo": "slide4.xml", "con_eduardo": "slide5.xml"}
HCF_TERMS_SLIDES = ["slide10.xml", "slide11.xml"]
HCF_FEE_SLIDE = "slide12.xml"

# Texto exacto (tal y como está en la plantilla) del bloque de garantía y
# de la firma, común a las dos plantillas — ver team.py para las opciones.
_WARRANTY_OLD = "dentro de los 3 meses siguientes a la fecha de contratación"
_DATE_PLACEHOLDER = "{dd/mm/aaaa}"
_FEE_OLD = "13%"

# ─── DONANTES PARA CLONAR "ÍNDICE" Y "CONTEXTO" ───────────────────────────
# Cada plantilla tiene ya una diapositiva con el "cascarón" que se necesita
# (logo, línea, texto de sección, título grande, subtítulo de una línea,
# número de página): en plantilla_hc.pptx es "Por qué Tessera" (slide5), en
# plantilla_hc_finance.pptx su equivalente (slide9, mismo patrón con ids
# distintos). Se clona dos veces, se le quitan las tarjetas propias de esa
# diapositiva (quedan solo el título y el subtítulo, reutilizado como
# cuerpo) y se reescribe el contenido.
_HC_DONOR = {
    "slide": "slide5.xml",
    "prune_ids": [286, 287, 289, 290, 292, 293, 4, 16, 7, 15, 19, 13, 14, 17, 18, 20, 21, 5, 6],
    "eyebrow_old": "NUESTRA FUERZA",
    "title_old": "Expertos en el flujo de talento",
    "body_old": "Procesos claros y compromisos que se cumplen: así movemos el talento.",
}
_HCF_DONOR = {
    "slide": "slide9.xml",
    "prune_ids": [185, 173, 174, 175, 176, 177, 183, 184, 186, 187, 188, 190, 191, 192, 193, 194, 195, 196, 197],
    "eyebrow_old": "NOSOTROS",
    "title_old": "Por qué Tessera",
    "body_old": "Control, rigor financiero y talento al servicio de la tesis de inversión.",
}

_INDEX_SLIDE_FILE = "slide90.xml"
_CONTEXT_SLIDE_FILE = "slide91.xml"


def _read(path: Path) -> str:
    return path.read_text("utf-8")


def _write(path: Path, content: str) -> None:
    path.write_text(content, "utf-8")


# ─── UTILIDADES DE PAQUETE (borrar / duplicar diapositivas) ───────────────
def _remove_slides(unpack: Path, slide_files: list[str]) -> None:
    """
    Quita por completo una o varias diapositivas del paquete .pptx: la
    referencia en <p:sldIdLst> de presentation.xml, la relación en
    presentation.xml.rels, el Override en [Content_Types].xml y el propio
    archivo de la diapositiva (con sus _rels). Es la forma correcta de
    borrar una diapositiva a nivel de paquete OOXML.
    """
    pres_path = unpack / "ppt" / "presentation.xml"
    pres_rels_path = unpack / "ppt" / "_rels" / "presentation.xml.rels"
    ct_path = unpack / "[Content_Types].xml"

    presentation_xml = _read(pres_path)
    pres_rels_xml = _read(pres_rels_path)
    content_types_xml = _read(ct_path)

    for slide_file in slide_files:
        target = f"slides/{slide_file}"
        rid = resolve_rid_for_target(pres_rels_xml, target)
        if rid:
            presentation_xml = remove_sld_id(presentation_xml, rid)
            pres_rels_xml = remove_relationship(pres_rels_xml, rid)
        content_types_xml = remove_content_type_override(content_types_xml, f"/ppt/slides/{slide_file}")

        slide_path = unpack / "ppt" / "slides" / slide_file
        slide_rels_path = unpack / "ppt" / "slides" / "_rels" / f"{slide_file}.rels"
        slide_path.unlink(missing_ok=True)
        slide_rels_path.unlink(missing_ok=True)

    _write(pres_path, presentation_xml)
    _write(pres_rels_path, pres_rels_xml)
    _write(ct_path, content_types_xml)


def _duplicate_slide(unpack: Path, source_slide_file: str, new_slide_file: str, insert_after_file: str) -> None:
    """
    Clona una diapositiva existente (XML + relaciones) y la inserta en el
    orden visual justo después de `insert_after_file`. No duplica media
    nueva: comparte las relaciones de la original (layout, imágenes
    decorativas) — con eso basta para clonar un "cascarón" ya con el
    estilo correcto, listo para que quien llame sustituya su contenido.
    """
    slides_dir = unpack / "ppt" / "slides"
    rels_dir = slides_dir / "_rels"

    shutil.copy(slides_dir / source_slide_file, slides_dir / new_slide_file)
    src_rels = rels_dir / f"{source_slide_file}.rels"
    if src_rels.exists():
        shutil.copy(src_rels, rels_dir / f"{new_slide_file}.rels")

    ct_path = unpack / "[Content_Types].xml"
    ct = _read(ct_path)
    override = (
        f'<Override PartName="/ppt/slides/{new_slide_file}" '
        f'ContentType="application/vnd.openxmlformats-officedocument.presentationml.slide+xml"/>'
    )
    _write(ct_path, ct.replace("</Types>", override + "</Types>"))

    pres_rels_path = unpack / "ppt" / "_rels" / "presentation.xml.rels"
    pres_rels = _read(pres_rels_path)
    existing_rids = [int(m) for m in re.findall(r'Id="rId(\d+)"', pres_rels)]
    new_rid = f"rId{max(existing_rids) + 1}"
    relationship = (
        f'<Relationship Id="{new_rid}" '
        f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide" '
        f'Target="slides/{new_slide_file}"/>'
    )
    pres_rels = pres_rels.replace("</Relationships>", relationship + "</Relationships>")
    _write(pres_rels_path, pres_rels)

    pres_path = unpack / "ppt" / "presentation.xml"
    pres = _read(pres_path)
    existing_ids = [int(m) for m in re.findall(r'<p:sldId\s+id="(\d+)"', pres)]
    new_id = max(existing_ids) + 1
    new_sldid = f'<p:sldId id="{new_id}" r:id="{new_rid}"/>'
    insert_after_rid = resolve_rid_for_target(pres_rels, f"slides/{insert_after_file}")
    anchor = re.search(r'<p:sldId\b[^>]*\br:id="%s"[^>]*/>' % re.escape(insert_after_rid), pres)
    pres = pres[:anchor.end()] + new_sldid + pres[anchor.end():]
    _write(pres_path, pres)


def _renumber_visible_pages(pptx_bytes: bytes) -> bytes:
    """
    Tras quitar o añadir diapositivas, el número de página visible en la
    esquina inferior derecha deja de ser correlativo. Se recorren las
    diapositivas ya ensambladas y se renumeran de forma secuencial las que
    tengan una forma con esa pinta (texto de 1-3 dígitos, en la esquina
    inferior derecha) — la portada y el cierre normalmente no llevan
    número, así que se saltan solas.
    """
    prs = Presentation(io.BytesIO(pptx_bytes))
    slide_w, slide_h = prs.slide_width, prs.slide_height
    visible = 0
    for slide in prs.slides:
        for shp in slide.shapes:
            if shp.left is None or not shp.has_text_frame:
                continue
            txt = shp.text_frame.text.strip()
            if not re.fullmatch(r"0?\d{1,2}", txt) or txt in ("2025", "2026"):
                continue
            if shp.left < slide_w * 0.6 or shp.top < slide_h * 0.8:
                continue
            visible += 1
            new_txt = f"{visible:02d}"
            tf = shp.text_frame
            first = True
            for p in tf.paragraphs:
                for r in p.runs:
                    if first:
                        r.text = new_txt
                        first = False
                    else:
                        r.text = ""
            break
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ─── ÍNDICE Y CONTEXTO (compartido) ────────────────────────────────────
def _build_info_slide(
    unpack: Path, donor: dict, new_slide_file: str, insert_after_file: str,
    eyebrow: str, title: str, body: str,
) -> None:
    _duplicate_slide(unpack, donor["slide"], new_slide_file, insert_after_file)
    path = unpack / "ppt" / "slides" / new_slide_file
    s = _read(path)
    for shape_id in donor["prune_ids"]:
        s = remove_shape_by_id(s, shape_id)
    s = replace_text_in_xml(s, donor["eyebrow_old"], eyebrow)
    s = replace_text_and_fit(s, donor["title_old"], title, role="title")
    s = replace_text_and_fit(s, donor["body_old"], body)
    s = enable_shrink_autofit(s)
    _write(path, s)


def _index_items(data: dict) -> list[str]:
    is_hc_finance = data.get("business_line") == BUSINESS_LINE_HC_FINANCE
    include_fee = data.get("deck_type") == "propuesta"
    if is_hc_finance:
        items = [
            "Contexto", "Tessera en una página", "Nuestro equipo",
            "Para el fondo", "Servicios Finance", "Servicios Human Capital",
            "Por qué Tessera",
        ]
        if include_fee:
            items.append("Honorarios")
        items += ["Confían en nosotros", "Contacto"]
    else:
        items = ["Contexto", "Nuestro equipo", "Dónde estamos", "Por qué Tessera"]
        items += [SERVICES_HC.get(s, s) for s in data.get("services", [])]
        if include_fee:
            items.append("Honorarios")
        items.append("Confían en nosotros")
    return items


def _build_index_and_context(data: dict, unpack: Path, donor: dict) -> None:
    items = _index_items(data)
    index_body = "\n".join(f"{i:02d}.  {name}" for i, name in enumerate(items, 1))
    _build_info_slide(
        unpack, donor, _INDEX_SLIDE_FILE, "slide1.xml",
        eyebrow="NOSOTROS", title="Índice", body=index_body,
    )

    try:
        context_body = generate_context_paragraph(data)
    except Exception:
        context_body = None
    context_body = context_body or (
        f"Tessera trabaja junto a {data.get('client_name', 'tu empresa')} para responder "
        f"a sus necesidades actuales con un equipo senior y un enfoque directo."
    )
    _build_info_slide(
        unpack, donor, _CONTEXT_SLIDE_FILE, _INDEX_SLIDE_FILE,
        eyebrow="NOSOTROS", title="Contexto", body=context_body,
    )


# ─── TÉRMINOS Y CONDICIONES + HONORARIOS (compartido) ─────────────────────
def _apply_terms_and_fee(
    unpack: Path, terms_slides: list[str], fee_slide: str,
    include_fee: bool, warranty_months: int, fee_rate: str | None,
) -> None:
    slides = unpack / "ppt" / "slides"

    if not include_fee:
        _remove_slides(unpack, terms_slides + [fee_slide])
        return

    warranty_new = _WARRANTY_OLD.replace("3 meses", f"{warranty_months} meses")
    today = datetime.date.today().strftime("%d/%m/%Y")

    for slide_file in terms_slides:
        path = slides / slide_file
        if not path.exists():
            continue
        s = _read(path)
        s = replace_text_in_xml(s, _WARRANTY_OLD, warranty_new)
        s = replace_text_in_xml(s, _DATE_PLACEHOLDER, today)
        _write(path, s)

    if fee_rate:
        path = slides / fee_slide
        if path.exists():
            _write(path, replace_text_in_xml(_read(path), _FEE_OLD, f"{fee_rate}%"))


# ─── HUMAN CAPITAL (solo) ──────────────────────────────────────────────
def _build_hc(data: dict, content: dict, unpack: Path) -> None:
    slides = unpack / "ppt" / "slides"
    r = replace_text_and_fit

    _build_index_and_context(data, unpack, _HC_DONOR)

    show_all_locations = data.get("locations", "all") != LOCATIONS_MADRID_ONLY
    drop_location = HC_LOCATION_SLIDES["madrid" if show_all_locations else "all"]
    _remove_slides(unpack, [drop_location])

    why = content.get("why_tessera", {})
    svcs = content.get("services", {})
    services = data.get("services", [])

    # ── Por qué Tessera ─────────────────────────────────────────────
    if why:
        s = _read(slides / "slide5.xml")
        if why.get("d1_title"): s = r(s, "Sin CVs al azar", why["d1_title"])
        if why.get("d1_body"):  s = r(s, "No enviamos el primer CV, buscamos a quien encaja de verdad. Candidatos filtrados sobre la mesa en 72 horas.", why["d1_body"])
        if why.get("d2_title"): s = r(s, "Sin pausas", why["d2_title"])
        if why.get("d2_body"):  s = r(s, "Tu servicio no parará.", why["d2_body"])
        if why.get("d3_title"): s = r(s, "Siempre contigo", why["d3_title"])
        if why.get("d3_body"):  s = r(s, "No desaparecemos tras la incorporación: cuidamos a la persona y al cliente.", why["d3_body"])
        _write(slides / "slide5.xml", s)

    # ── Headhunting ──────────────────────────────────────────────────
    if "headhunting" in services and svcs.get("headhunting"):
        hh = svcs["headhunting"]
        s = _read(slides / "slide6.xml")
        if hh.get("why_col_title1"): s = r(s, "VELOCIDAD", hh["why_col_title1"])
        if hh.get("why_col_body1"):  s = r(s, "Disponemos de nuestra propia metodología, la cual nos permite presentar candidatos a tiempo.", hh["why_col_body1"])
        if hh.get("why_col_title2"): s = r(s, "ESPECIALIZACIÓN", hh["why_col_title2"])
        if hh.get("why_col_body2"):  s = r(s, "Conocemos tu sector y cómo se construyen sus equipos.", hh["why_col_body2"])
        if hh.get("why_col_title3"): s = r(s, "ÉXITO COMPARTIDO", hh["why_col_title3"])
        if hh.get("why_col_body3"):  s = r(s, "El equipo que lleva tu negocio al siguiente nivel.", hh["why_col_body3"])
        if hh.get("benefits_body1"): s = r(s, "Nos ocupamos de todo. Tú solo conoces a los mejores.", hh["benefits_body1"])
        if hh.get("benefits_body2"): s = r(s, "Candidatos que no solo encajan: impulsan tu crecimiento.", hh["benefits_body2"])
        if hh.get("benefits_body3"): s = r(s, "Tu éxito es también el nuestro.", hh["benefits_body3"])
        if hh.get("how_body3"): s = r(s, "Habilidades, actitud y encaje cultural.", hh["how_body3"])
        if hh.get("how_body4"): s = r(s, "Únicamente candidatos que realmente suman.", hh["how_body4"])
        if hh.get("how_body5"): s = r(s, "Contigo también después de la incorporación.", hh["how_body5"])
        _write(slides / "slide6.xml", s)

    # ── Outsourcing ──────────────────────────────────────────────────
    if "outsourcing" in services and svcs.get("outsourcing"):
        outs = svcs["outsourcing"]
        s = _read(slides / "slide7.xml")
        if outs.get("headline"): s = r(s, "Externalización que sí funciona: pagas por trabajo real, sin papeleos.", outs["headline"])
        if outs.get("body"):     s = r(s, "Contratar cuesta más de lo que parece. Con el modelo Time & Material ajustas el equipo a tu actividad real y nosotros nos encargamos de toda la gestión.", outs["body"])
        if outs.get("card1_body"): s = r(s, "Pagas solo por horas reales de trabajo.", outs["card1_body"])
        if outs.get("card2_body"): s = r(s, "Cubrimos bajas y vacaciones.", outs["card2_body"])
        if outs.get("card3_body"): s = r(s, "Nóminas y trámites, a cargo nuestro.", outs["card3_body"])
        if outs.get("card4_body"): s = r(s, "Escalas el equipo cuando quieras.", outs["card4_body"])
        _write(slides / "slide7.xml", s)

    # ── Consultoría Salarial ─────────────────────────────────────────
    if "salary" in services and svcs.get("salary"):
        sal = svcs["salary"]
        s = _read(slides / "slide8.xml")
        if sal.get("headline"): s = r(s, "Inteligencia retributiva para decisiones que importan.", sal["headline"])
        if sal.get("body"):     s = r(s, "Acompañamos a organizaciones en el diseño de estructuras salariales coherentes, alineadas con el mercado y preparadas para el nuevo marco regulatorio europeo.", sal["body"])
        _write(slides / "slide8.xml", s)

    # ── Formación ────────────────────────────────────────────────────
    if "formacion" in services and svcs.get("formacion"):
        form = svcs["formacion"]
        s = _read(slides / "slide9.xml")
        if form.get("body"): s = r(s, "Potenciamos el talento interno: consultoría en formación orientada a procesos y resultados ágiles.", form["body"])
        _write(slides / "slide9.xml", s)

    # ── Diapositivas de servicio no seleccionadas: se quitan ─────────
    to_remove = [fn for key, fn in HC_SERVICE_SLIDES.items() if key not in services]
    if to_remove:
        _remove_slides(unpack, to_remove)

    # ── Autofit dinámico en las diapositivas con texto variable ──────
    for name in HC_SLIDES_DYNAMIC_TEXT:
        path = slides / name
        if path.exists():
            _write(path, enable_shrink_autofit(_read(path)))

    _apply_terms_and_fee(
        unpack, HC_TERMS_SLIDES, HC_FEE_SLIDE,
        include_fee=data.get("deck_type") == "propuesta",
        warranty_months=data.get("warranty_months", 3),
        fee_rate=data.get("fee_rate"),
    )


# ─── HUMAN CAPITAL + FINANCE ────────────────────────────────────────────
def _build_hc_finance(data: dict, content: dict, unpack: Path) -> None:
    _build_index_and_context(data, unpack, _HCF_DONOR)

    show_all_locations = data.get("locations", "all") == "all"
    keep_location = HCF_LOCATION_SLIDES["all" if show_all_locations else "madrid"]
    drop_location = HCF_LOCATION_SLIDES["madrid" if show_all_locations else "all"]

    include_eduardo = bool(data.get("include_eduardo"))
    keep_team = HCF_TEAM_SLIDES["con_eduardo" if include_eduardo else "no_eduardo"]
    drop_team = HCF_TEAM_SLIDES["no_eduardo" if include_eduardo else "con_eduardo"]

    _remove_slides(unpack, [drop_location, drop_team])

    _apply_terms_and_fee(
        unpack, HCF_TERMS_SLIDES, HCF_FEE_SLIDE,
        include_fee=data.get("deck_type") == "propuesta",
        warranty_months=data.get("warranty_months", 3),
        fee_rate=data.get("fee_rate"),
    )


# ─── PUNTO DE ENTRADA ────────────────────────────────────────────────────
def build_pptx(data: dict, content: dict, logo_bytes: bytes | None = None) -> bytes:
    business_line = data.get("business_line", "hc")
    plantilla = PLANTILLA_HC_FINANCE if business_line == BUSINESS_LINE_HC_FINANCE else PLANTILLA_HC

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        work = tmp / "work.pptx"
        shutil.copy(str(plantilla), str(work))

        unpack = tmp / "unpacked"
        with zipfile.ZipFile(work, 'r') as z:
            z.extractall(unpack)

        if business_line == BUSINESS_LINE_HC_FINANCE:
            _build_hc_finance(data, content, unpack)
        else:
            _build_hc(data, content, unpack)

        out = tmp / "output.pptx"
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
            for fp in sorted(unpack.rglob('*')):
                if fp.is_file():
                    zout.write(fp, fp.relative_to(unpack))

        return _renumber_visible_pages(out.read_bytes())
